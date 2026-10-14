"""Справочник препаратов: стартовый набор + выгрузка ГРЛС.

ГРЛС — Государственный реестр лекарственных средств (grls.rosminzdrav.ru).
Реестр выгружается архивом с таблицами Excel. Заголовки колонок в разных
выгрузках отличаются, поэтому загрузчик находит строку заголовка и колонки
по ключевым словам («торгов…», «МНН/непатентован…», «АТХ»), а если формат
не распознан — принимает явное сопоставление `columns={...}`.

Справочник отвечает на три вопроса:

* какое МНН у торгового наименования («Нурофен» → ибупрофен) — это ключ факта в памяти;
* из каких действующих веществ состоит комбинированный препарат;
* какие у препарата коды АТХ — по ним правила безопасности относят препарат к классам.

Записи стартового набора имеют приоритет: если торговое наименование из ГРЛС
совпадает с синонимом из стартового набора, остаётся стартовая запись, поэтому
ключи памяти в тестах и в golden set не меняются при подключении ГРЛС.
"""
from __future__ import annotations

import csv
import gzip
import io
import json
import re
import zipfile
from collections import Counter
from dataclasses import asdict, dataclass, field, replace
from itertools import islice
from pathlib import Path
from typing import Any, Iterable, Iterator, Sequence

from . import reference_data as ref
from .lexicon import LexEntry, Lexicon, NormResult, clean_name, fold

# =========================================================================== разбор ячеек

_CYR_LOOKALIKE = str.maketrans("АВСЕНКМОРТХУ", "ABCEHKMOPTXY")
_ATC_RE = re.compile(r"(?<![A-Z0-9])([A-Z]\d{2}(?:[A-Z](?:[A-Z](?:\d{2})?)?)?)(?![A-Z0-9])")
ATC_GROUPS = frozenset("ABCDGHJLMNPRSV")


def parse_atc(value: Any) -> tuple[str, ...]:
    """Коды АТХ из ячейки: «N02BE01», «N02BE01 Парацетамол», «М01АЕ01» (кириллица вместо латиницы).

    Коды короче 4 символов («B12» в названии витамина) принимаются, только если ячейка из него и состоит.
    """
    if value is None:
        return ()
    s = str(value).upper().translate(_CYR_LOOKALIKE)
    codes = [c for c in _ATC_RE.findall(s) if c[0] in ATC_GROUPS]
    long = [c for c in codes if len(c) >= 4]
    if long:
        return tuple(dict.fromkeys(long))
    if len(codes) == 1 and s.strip() == codes[0]:
        return (codes[0],)
    return ()


_NO_INN = frozenset({"~", "-", "—", "–", "нет", "отсутствует", "не применимо", "не указано", "не имеется"})
_INN_SPLIT = re.compile(r"\s*(?:\+|;|\n)\s*")
_HAS_LETTER = re.compile(r"[a-zа-я]")

# МНН в ГРЛС, которые в стартовом наборе записаны иначе
INN_SYNONYMS: dict[str, str] = {
    "смектит диоктаэдрический": "диосмектит",
    "уголь активированный": "активированный уголь",
    "активированный уголь": "активированный уголь",
    "левотироксин": "левотироксин натрия",
}
# Одно МНН в ГРЛС — два вещества в стартовом наборе
INN_SPLIT_SYNONYMS: dict[str, tuple[str, ...]] = {
    "калия и магния аспарагинат": ("калия аспарагинат", "магния аспарагинат"),
}


def parse_inn(value: Any) -> tuple[str, ...]:
    """Действующие вещества из ячейки МНН: «Амоксициллин+[Клавулановая кислота]» → (амоксициллин, клавулановая кислота)."""
    if value is None:
        return ()
    s = fold(str(value))
    for ch in "[]*®™":
        s = s.replace(ch, " ")
    s = re.sub(r"\([^)]*\)", " ", s)  # пояснения в скобках
    parts: list[str] = []
    for p in _INN_SPLIT.split(s):
        p = re.sub(r"\s+", " ", p).strip(" .,;:")
        if p and p not in _NO_INN and _HAS_LETTER.search(p):
            parts.extend(INN_SPLIT_SYNONYMS.get(p, (INN_SYNONYMS.get(p, p),)))
    return tuple(dict.fromkeys(parts))


def clean_trade_name(value: Any) -> str:
    s = str(value or "")
    for ch in "®™©":
        s = s.replace(ch, "")
    s = re.sub(r"[«»\"“”„]", " ", s)
    return re.sub(r"\s+", " ", s).strip(" .,;")


# =========================================================================== загрузка ГРЛС


# Путь введения по названию лекарственной формы. Нужен, потому что в выгрузке ГРЛС нет кодов АТХ,
# а правило «НПВП при язве» не должно срабатывать на гель для наружного применения.
ROUTES = ("systemic", "skin", "eye", "ear", "nasal", "vaginal", "mucosal", "inhal")
_ROUTE_PATTERNS: tuple[tuple[str, re.Pattern[str]], ...] = (
    ("systemic", re.compile(r"внутрь|инъекц|инфуз|внутривен|внутримышеч|подкож|ректальн|подъязычн|"
                            r"трансдермальн|защечн")),
    ("eye", re.compile(r"глазн|офтальм|конъюнктив")),
    ("ear", re.compile(r"ушн")),
    ("nasal", re.compile(r"назальн|в нос|интраназ")),
    ("vaginal", re.compile(r"вагинальн")),
    ("inhal", re.compile(r"ингаляц")),
    ("mucosal", re.compile(r"для рассасывания|пастилк|леденц|для местного|полоскани|стоматолог|зубн|"
                           r"для приготовления раствора для местного")),
    ("skin", re.compile(r"наружн|накожн|мазь|крем|гель|лосьон|шампун|линимент|пена|пластыр|присыпк|"
                        r"лак для ногтей|паста")),
    ("systemic", re.compile(r"таблет|капсул|драже|сироп|суспенз|гранул|порошок|раствор|капли|эликсир|"
                            r"настойк|суппозит|лиофилизат|концентрат|эмульс")),
)
_OTC_RE = re.compile(r"без\s+рецепта")
_RX_RE = re.compile(r"по\s+рецепту")


def form_route(form_name: str) -> str | None:
    """«гель для наружного применения» → skin; «таблетки, покрытые оболочкой» → systemic."""
    f = fold(form_name)
    for route, pattern in _ROUTE_PATTERNS:
        if pattern.search(f):
            return route
    return None


def parse_forms(value: Any) -> tuple[tuple[str, ...], bool | None]:
    """Пути введения и порядок отпуска из ячейки «Формы выпуска».

    Ячейка ГРЛС: «таблетки, 200 мг, 10 шт. - блистеры - пачки картонные - Без рецепта; гель для наружного
    применения, 5 %, 50 г - тубы - По рецепту; …». Отпуск: True — есть упаковка без рецепта, False — все
    упаковки по рецепту, None — не указано.
    """
    if value is None:
        return (), None
    text = fold(str(value))
    routes: list[str] = []
    for seg in text.split(";"):
        seg = seg.strip()
        if not seg:
            continue
        route = form_route(seg.split(",")[0])
        if route and route not in routes:
            routes.append(route)
    otc = True if _OTC_RE.search(text) else (False if _RX_RE.search(text) else None)
    return tuple(routes), otc


@dataclass(frozen=True)
class Product:
    """Одна строка реестра: торговое наименование и его состав."""

    trade_name: str
    components: tuple[str, ...]
    atc: tuple[str, ...] = ()
    reg_number: str | None = None
    form: str | None = None
    routes: tuple[str, ...] = ()   # пути введения из лекарственных форм
    otc: bool | None = None        # есть ли упаковка без рецепта


# Роль колонки → варианты набора ключевых слов (все слова варианта должны быть в заголовке).
# Английские варианты — заголовки набора открытых данных Минздрава (CSV).
COLUMN_KEYWORDS: dict[str, tuple[tuple[str, ...], ...]] = {
    "trade_name": (("торгов",), ("tradename",)),
    "inn": (("мнн",), ("непатентован",), ("группировочн",), ("internationalname",)),
    "atc": (("атх",), ("анатомо",)),
    "reg_number": (("регистрацион", "номер"), ("regnumber",)),
    "form": (("форм", "выпуск"), ("лекарствен", "форм"), ("formrelease",)),
    "annulled": (("аннулир",), ("исключ",), ("cancellationdate",)),
}
REQUIRED_COLUMNS = ("trade_name", "inn")
EXCLUDED_TABLE_RE = re.compile(r"исключ|аннулир")


def _cell_text(v: Any) -> str:
    return "" if v is None else re.sub(r"\s+", " ", fold(str(v)))


def detect_columns(header: Sequence[Any]) -> dict[str, int]:
    """Роль → номер колонки по ключевым словам заголовка."""
    texts = [_cell_text(v) for v in header]
    found: dict[str, int] = {}
    for role, variants in COLUMN_KEYWORDS.items():
        for i, text in enumerate(texts):
            if i in found.values() or not text:
                continue
            if any(all(k in text for k in variant) for variant in variants):
                found[role] = i
                break
    return found


def _manual_columns(header: Sequence[Any], columns: dict[str, str | int]) -> dict[str, int] | None:
    texts = [_cell_text(v) for v in header]
    out: dict[str, int] = {}
    for role, spec in columns.items():
        if isinstance(spec, int):
            out[role] = spec
            continue
        want = _cell_text(spec)
        idx = next((i for i, t in enumerate(texts) if t == want), None)
        if idx is None:
            idx = next((i for i, t in enumerate(texts) if want and want in t), None)
        if idx is None:
            return None
        out[role] = idx
    return out


def find_header(rows: Sequence[Sequence[Any]], columns: dict[str, str | int] | None = None,
                search_rows: int = 30) -> tuple[int, int, dict[str, int]] | None:
    """(первая строка заголовка, число строк заголовка, колонки) или None.

    Над заголовком в выгрузке бывают строки с названием реестра и датой; заголовок
    бывает разбит на две строки (объединённые ячейки) — тогда строки склеиваются.
    """
    if columns and all(isinstance(v, int) for v in columns.values()):
        # Только номера колонок: заголовок — первая непустая строка, если не найден по словам
        for r, row in enumerate(rows[:search_rows]):
            if any(_cell_text(v) for v in row):
                auto = detect_columns(row)
                if all(k in auto for k in REQUIRED_COLUMNS):
                    return r, 1, {**auto, **columns}
        return 0, 1, dict(columns)
    def match(header: Sequence[Any]) -> dict[str, int] | None:
        manual = _manual_columns(header, columns) if columns else None
        if columns and manual is None:
            return None
        cols = {**detect_columns(header), **(manual or {})}
        return cols if all(k in cols for k in REQUIRED_COLUMNS) else None

    def merged(r: int) -> list[str]:
        top, nxt = list(rows[r]), list(rows[r + 1])
        width = max(len(top), len(nxt))
        top += [None] * (width - len(top))
        nxt += [None] * (width - len(nxt))
        return [" ".join(_cell_text(a) for a in pair if a is not None) for pair in zip(top, nxt)]

    limit = min(search_rows, len(rows))
    for r in range(limit):
        has_next = r + 1 < len(rows)
        cols = match(rows[r])
        if cols:
            # «Торговое | Международное непатентованное | Код» + «наименование | наименование | АТХ»:
            # вторая строка — продолжение заголовка, если вместе с ней находится больше колонок
            if has_next and not match(rows[r + 1]):
                both = match(merged(r))
                if both and len(both) > len(cols):
                    return r, 2, both
            return r, 1, cols
        # Начало заголовка в строке r, окончание — в следующей (объединённые ячейки)
        if has_next and detect_columns(rows[r]) and not match(rows[r + 1]):
            both = match(merged(r))
            if both:
                return r, 2, both
    return None


@dataclass
class TableReport:
    name: str
    used: bool
    reason: str = ""
    header_row: int | None = None
    columns: dict[str, int] = field(default_factory=dict)
    header: list[str] = field(default_factory=list)
    rows: int = 0
    products: int = 0
    skipped: dict[str, int] = field(default_factory=dict)


@dataclass
class LoadReport:
    path: str
    tables: list[TableReport] = field(default_factory=list)

    @property
    def products(self) -> int:
        return sum(t.products for t in self.tables)

    def as_dict(self) -> dict[str, Any]:
        return {"path": self.path, "products": self.products, "tables": [asdict(t) for t in self.tables]}


def _decode_zip_name(info: zipfile.ZipInfo) -> str:
    if info.flag_bits & 0x800:
        return info.filename
    try:
        return info.filename.encode("cp437").decode("cp866")
    except (UnicodeEncodeError, UnicodeDecodeError):
        return info.filename


def _csv_rows(data: bytes) -> Iterator[list[str]]:
    for enc in ("utf-8-sig", "cp1251"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = data.decode("utf-8", errors="replace")
    head = text[:5000]
    delim = max((";", "\t", ","), key=head.count)
    yield from csv.reader(io.StringIO(text), delimiter=delim)


def _tables_from_bytes(name: str, ext: str, data: bytes) -> Iterator[tuple[str, Iterator[Sequence[Any]]]]:
    if ext == ".xlsx":
        import openpyxl  # requirements-data.txt

        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for ws in wb.worksheets:
            yield f"{name}:{ws.title}", ws.iter_rows(values_only=True)
    elif ext == ".xls":
        try:
            import xlrd  # requirements-data.txt
        except ImportError as exc:  # pragma: no cover - зависит от окружения
            raise ImportError("для .xls нужен пакет xlrd: pip install xlrd") from exc
        book = xlrd.open_workbook(file_contents=data)
        for sheet in book.sheets():
            yield f"{name}:{sheet.name}", (sheet.row_values(i) for i in range(sheet.nrows))
    elif ext in (".csv", ".txt"):
        yield name, _csv_rows(data)


GRLS_EXTENSIONS = (".zip", ".xlsx", ".xls", ".csv")


def grls_files(path: str | Path | Iterable[str | Path]) -> list[Path]:
    """Файлы выгрузки: один файл, список файлов или папка (все .zip/.xlsx/.xls/.csv в ней).

    ГРЛС выгружается набором файлов по статусу записи: «Действующий», «Изменённый», «Истёкший»,
    «Исключённый»… — папку можно передать целиком.
    """
    if isinstance(path, (str, Path)):
        p = Path(path)
        if p.is_dir():
            return sorted(x for x in p.iterdir() if x.is_file() and x.suffix.lower() in GRLS_EXTENSIONS)
        return [p]
    return [Path(x) for x in path]


def iter_tables(path: str | Path) -> Iterator[tuple[str, Iterator[Sequence[Any]]]]:
    """Таблицы из файла выгрузки: .zip (с xls/xlsx/csv внутри), .xlsx, .xls, .csv."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".zip":
        with zipfile.ZipFile(path) as zf:
            for info in zf.infolist():
                name = _decode_zip_name(info)
                inner = Path(name).suffix.lower()
                if inner in (".xlsx", ".xls", ".csv", ".txt"):
                    yield from _tables_from_bytes(name, inner, zf.read(info))
    else:
        yield from _tables_from_bytes(path.name, ext, path.read_bytes())


def load_grls(
    path: str | Path | Iterable[str | Path],
    *,
    columns: dict[str, str | int] | None = None,
    include_excluded: bool = False,
    search_rows: int = 30,
) -> tuple[list[Product], LoadReport]:
    """Читает выгрузку реестра (файл, список файлов или папку). Возвращает строки-препараты и отчёт
    о том, что и почему пропущено.

    columns — явное сопоставление роли и колонки: {"trade_name": "Торговое наименование", "inn": 8, "atc": "АТХ"}.
    Файлы и таблицы исключённых записей и строки с датой аннулирования пропускаются, если не include_excluded.
    """
    files = grls_files(path)
    report = LoadReport(", ".join(str(f) for f in files))
    products: list[Product] = []
    seen: set[str] = set()
    tables = ((name, rows) for f in files for name, rows in iter_tables(f))
    for name, rows_iter in tables:
        tr = TableReport(name=name, used=False)
        report.tables.append(tr)
        # Архив папки с Google Drive содержит те же файлы, что лежат рядом: повтор не читаем
        file_part, _, sheet = name.partition(":")
        key = f"{Path(file_part).name}:{sheet}"
        if key in seen:
            tr.reason = "повтор: такая таблица уже прочитана"
            continue
        seen.add(key)
        if not include_excluded and EXCLUDED_TABLE_RE.search(fold(name)):
            tr.reason = "таблица исключённых записей"
            continue
        head = list(islice(rows_iter, search_rows + 1))
        found = find_header(head, columns, search_rows)
        if found is None:
            tr.reason = "не найдены колонки торгового наименования и МНН"
            tr.header = [_cell_text(v) for v in (head[0] if head else [])][:30]
            continue
        start, n_header, cols = found
        tr.used, tr.header_row, tr.columns = True, start, cols
        tr.header = [_cell_text(v) for v in head[start]][:30]
        skipped: Counter = Counter()

        def cell(row: Sequence[Any], role: str) -> Any:
            i = cols.get(role)
            return row[i] if i is not None and i < len(row) else None

        def rows() -> Iterator[Sequence[Any]]:
            yield from head[start + n_header:]
            yield from rows_iter

        for row in rows():
            if not row or not any(v not in (None, "") for v in row):
                continue
            tr.rows += 1
            trade = clean_trade_name(cell(row, "trade_name"))
            if not trade or not _HAS_LETTER.search(fold(trade)):
                skipped["нет торгового наименования"] += 1
                continue
            if not include_excluded and _cell_text(cell(row, "annulled")):
                skipped["аннулировано"] += 1
                continue
            routes, otc = parse_forms(cell(row, "form"))
            products.append(Product(
                trade_name=trade,
                components=parse_inn(cell(row, "inn")),
                atc=parse_atc(cell(row, "atc")),
                reg_number=(str(cell(row, "reg_number")).strip() or None) if cell(row, "reg_number") else None,
                form=(re.sub(r"\s+", " ", str(cell(row, "form"))).strip()[:200] or None)
                if cell(row, "form") else None,
                routes=routes, otc=otc,
            ))
            tr.products += 1
        tr.skipped = dict(skipped)
    return products, report


# =========================================================================== справочник


@dataclass(frozen=True)
class DrugInfo:
    """Препарат после сопоставления со справочником."""

    raw: str
    canonical: str
    components: tuple[str, ...]
    atc: tuple[str, ...]
    matched: bool
    method: str = "exact"
    routes: tuple[str, ...] = ()   # пути введения (известны из ГРЛС)
    otc: bool | None = None        # True — есть безрецептурные формы, False — только по рецепту, None — неизвестно

    @property
    def label(self) -> str:
        """«Нурофен (ибупрофен)», если назван торговым наименованием, иначе МНН."""
        raw = re.sub(r"\s+", " ", self.raw.strip())
        if not raw or clean_name(raw) == clean_name(self.canonical):
            return self.canonical
        return f"{raw[:1].upper()}{raw[1:]} ({self.canonical})"


@dataclass
class _Entry:
    canonical: str
    components: tuple[str, ...]
    aliases: list[str] = field(default_factory=list)
    atc: set[str] = field(default_factory=set)
    classes: tuple[str, ...] = ()
    sources: set[str] = field(default_factory=set)
    products: int = 0
    routes: set[str] = field(default_factory=set)
    otc: set[bool] = field(default_factory=set)


# Торговые наименования, совпадающие с обычными словами: в речи покупателя это не препарат.
# Ищутся в тексте только записи стартового набора и «отличимые» наименования ГРЛС (см. mention_keys).
MENTION_STOPWORDS = frozenset("""
актив баланс беби бэби вита витамин витамины детский доктор здоровье иммунитет комфорт лайт мама макс
норма нео плюс сила стоп форте экспресс эффект вода сода соль сахар масло мята ромашка шалфей календула
подорожник алоэ солодка чай кофе глюкоза кислород азот спирт вазелин глицерин парафин тальк сера
йод мед мёд крем гель мазь капли спрей сироп раствор таблетки капсулы порошок свечи пластырь бинт
""".split())


# Версия правил сборки справочника: кэш, собранный другой версией, не загружается
CATALOG_VERSION = 2

# Наименование, которое начинается с предлога, в речи почти всегда не препарат: «что выпить от кашля»
MENTION_PREPOSITIONS = frozenset("от для при против с со на в во по из без к о об".split())


class DrugCatalog:
    """Справочник препаратов: МНН, торговые наименования, состав, коды АТХ."""

    def __init__(self, seed: Iterable[LexEntry] | None = None,
                 alias_atc: dict[str, tuple[str, ...]] | None = None):
        self._entries: dict[str, _Entry] = {}
        self._by_components: dict[frozenset[str], str] = {}
        self._alias_owner: dict[str, str] = {}
        self._alias_source: dict[str, str] = {}
        self._ambiguous: dict[str, set[str]] = {}
        self._alias_atc: dict[str, set[str]] = {}
        self._alias_routes: dict[str, set[str]] = {}
        self._alias_otc: dict[str, set[bool]] = {}
        self._no_inn: set[str] = set()                       # записи, ключ которых — торговое наименование
        self._alias_counts: dict[str, Counter] = {}          # наименование → сколько строк ГРЛС с каждым составом
        self._alias_by: dict[str, dict[str, tuple[set, set, set]]] = {}  # наименование → состав → АТХ, формы, отпуск
        self._resolved: dict[str, str] = {}                  # неоднозначное наименование → выбранный состав
        self._alias_union: dict[str, tuple[str, ...]] = {}   # и все вещества из всех его составов
        self._seed_conflicts: dict[str, set[str]] = {}       # синоним стартового набора → состав по ГРЛС
        self._lexicon: Lexicon | None = None
        self.stats: Counter = Counter()
        for e in (ref.DRUGS if seed is None else seed):
            self._add_seed(e)
        for alias, codes in (ref.DRUG_ALIAS_ATC if alias_atc is None else alias_atc).items():
            self._alias_atc.setdefault(clean_name(alias), set()).update(codes)

    # ------------------------------------------------------------------ наполнение
    def _entry(self, canonical: str, components: tuple[str, ...], source: str) -> _Entry:
        entry = self._entries.get(canonical)
        if entry is None:
            entry = _Entry(canonical, components)
            self._entries[canonical] = entry
            self._by_components.setdefault(frozenset(components), canonical)
            self._register_alias(canonical, canonical, source)
        entry.sources.add(source)
        return entry

    def _register_alias(self, alias: str, canonical: str, source: str, atc: Iterable[str] = (),
                        product: Product | None = None) -> None:
        key = clean_name(alias)
        if not key:
            return
        attrs: tuple[set, set, set] = (set(atc), set(product.routes) if product else set(),
                                       {product.otc} if product is not None and product.otc is not None else set())
        if source != "seed":
            self._alias_counts.setdefault(key, Counter())[canonical] += 1
            # АТХ, формы и отпуск торгового наименования — по каждому составу: пригодятся, если
            # наименование окажется неоднозначным и будет разрешено по частоте
            by = self._alias_by.setdefault(key, {}).setdefault(canonical, (set(), set(), set()))
            for dst, src in zip(by, attrs):
                dst |= src
        owner = self._alias_owner.get(key)
        if (owner is not None and owner != canonical and owner in self._no_inn and canonical not in self._no_inn
                and clean_name(owner) == key and self._alias_source.get(key) != "seed"):
            # Сначала встретилась строка без МНН («Ингалипт»), теперь — с составом: запись без МНН сливается
            self._merge_entry(owner, canonical)
            self.stats["no_inn_merged"] += 1
            owner = canonical
        if owner is None:
            self._alias_owner[key] = canonical
            self._alias_source[key] = source
            self._entries[canonical].aliases.append(alias)
            owner = canonical
        elif owner != canonical:
            if self._alias_source[key] == "seed":
                self.stats["seed_priority"] += 1  # стартовая запись важнее ГРЛС
                self._seed_conflicts.setdefault(key, set()).add(canonical)
            elif clean_name(owner) == key:
                self.stats["inn_name_priority"] += 1  # «Парацетамол» — это МНН, а не чужое торговое наименование
            else:
                self._ambiguous.setdefault(key, {owner}).add(canonical)
            return
        if source != "seed":
            self._add_alias_attrs(key, attrs)
        self._lexicon = None

    def _add_alias_attrs(self, key: str, attrs: tuple[set, set, set]) -> None:
        for mapping, values in zip((self._alias_atc, self._alias_routes, self._alias_otc), attrs):
            if values:
                mapping.setdefault(key, set()).update(values)

    def _add_seed(self, e: LexEntry) -> None:
        components = e.components or (e.canonical,)
        entry = self._entry(e.canonical, components, "seed")
        entry.atc.update(e.atc)
        entry.classes = e.classes
        for alias in e.aliases:
            self._register_alias(alias, e.canonical, "seed")

    def add_product(self, p: Product, source: str = "grls") -> str:
        """Добавляет строку реестра. Возвращает канонический ключ (МНН или состав комбинированного)."""
        if p.components:
            # Синонимы МНН — и для строк, собранных не через parse_inn
            comps = tuple(dict.fromkeys(x for c in p.components
                                        for x in INN_SPLIT_SYNONYMS.get(c, (INN_SYNONYMS.get(c, c),))))
            canonical = self._by_components.get(frozenset(comps)) or " + ".join(comps)
            components = self._entries[canonical].components if canonical in self._entries else comps
        else:
            # Нет МНН (фитопрепараты, гомеопатия): ключом остаётся торговое наименование.
            # Если это наименование уже известно с составом («Алмагель» в стартовом наборе или
            # в другой строке ГРЛС), строка относится к той записи
            key = clean_name(p.trade_name) or fold(p.trade_name)
            key = INN_SYNONYMS.get(key, key)
            owner = self._alias_owner.get(key)
            canonical = owner if owner in self._entries and owner not in self._no_inn else key
            components = self._entries[canonical].components if canonical in self._entries else (canonical,)
            self.stats["no_inn"] += 1
            if canonical not in self._entries:
                self._no_inn.add(canonical)
        if p.components:
            self._no_inn.discard(canonical)
        entry = self._entry(canonical, components, source)
        entry.atc.update(p.atc)
        entry.routes.update(p.routes)
        if p.otc is not None:
            entry.otc.add(p.otc)
        entry.products += 1
        self.stats["products"] += 1
        self._register_alias(p.trade_name, canonical, source, p.atc, p)
        return canonical

    def add_products(self, products: Iterable[Product], source: str = "grls") -> int:
        n = 0
        for p in products:
            self.add_product(p, source)
            n += 1
        self.merge_inn_variants()
        self.resolve_ambiguous()
        self._lexicon = None
        return n

    def resolve_ambiguous(self) -> int:
        """Наименование с разным составом в разных строках ГРЛС → самый частый состав.

        Записи без МНН (ключ — само наименование) уступают записям с МНН. При равенстве частот
        наименование остаётся нераспознанным. Для проверки безопасности у наименования остаются
        все вещества из всех его составов: «Колдрекс» встречается и с фенилэфрином, и с хлорфенамином.
        """
        resolved = 0
        for key, cands in self._ambiguous.items():
            counts = self._alias_counts.get(key, Counter())
            names = [c for c in cands if c in self._entries]
            with_inn = [c for c in names if c not in self._no_inn]
            pool = with_inn or names
            ranked = sorted(pool, key=lambda c: (-counts.get(c, 0), c))
            if not ranked or (len(ranked) > 1 and counts.get(ranked[0], 0) == counts.get(ranked[1], 0)):
                self._resolved.pop(key, None)
                continue
            winner = ranked[0]
            self._resolved[key] = winner
            if self._alias_owner.get(key) != winner:
                self._alias_owner[key] = winner
                alias = next((a for c in names for a in self._entries[c].aliases if clean_name(a) == key), key)
                if all(clean_name(a) != key for a in self._entries[winner].aliases):
                    self._entries[winner].aliases.append(alias)
            for c in ranked:  # АТХ, формы и отпуск самого наименования — по всем его составам
                self._add_alias_attrs(key, self._alias_by.get(key, {}).get(c, (set(), set(), set())))
            union: list[str] = []
            for c in [winner, *ranked[1:]]:
                for comp in self._entries[c].components:
                    if comp not in union:
                        union.append(comp)
            self._alias_union[key] = tuple(union)
            resolved += 1
        self._lexicon = None
        return resolved

    def _usable(self, key: str, canonical: str | None = None) -> bool:
        """Наименование однозначно (или разрешено по частоте) и принадлежит записи canonical."""
        if key not in self._ambiguous:
            return True
        return key in self._resolved and (canonical is None or self._resolved[key] == canonical)

    _ENDINGS = ("ами", "ями", "ов", "ом", "ем", "ой", "ей", "ам", "ах", "а", "я", "у", "ю", "ы", "и", "е")

    def _inn_variant(self, name: str) -> str | None:
        """«парацетамола» → «парацетамол», если такое МНН есть в справочнике."""
        words = name.split()
        if not words:
            return None
        last = words[-1]
        for ending in self._ENDINGS:
            if last.endswith(ending) and len(last) - len(ending) >= 4:
                stem = last[: -len(ending)]
                for cand in (stem, stem + "а", stem + "я", stem + "ь"):
                    target = " ".join(words[:-1] + [cand])
                    if target != name and target in self._entries and target not in self._no_inn:
                        return target
        return None

    def merge_inn_variants(self) -> int:
        """Записи без МНН, у которых торговое наименование — МНН в другом падеже, сливаются с МНН.

        В ГРЛС встречаются препараты без МНН с торговым наименованием вроде «Парацетамола …»:
        без слияния слово «парацетамола» в ответе модели считалось бы отдельным препаратом.
        """
        merged = 0
        for name in sorted(self._no_inn):
            entry = self._entries.get(name)
            target = self._inn_variant(name)
            if entry is None or target is None or "seed" in entry.sources:
                continue
            self._merge_entry(name, target)
            merged += 1
        self.stats["merged_inn_variants"] += merged
        return merged

    def _merge_entry(self, src: str, dst: str) -> None:
        """Запись src (без МНН) становится частью dst: наименования, формы, отпуск, счётчики."""
        entry = self._entries.pop(src)
        target = self._entries[dst]
        for alias in entry.aliases:
            key = clean_name(alias)
            if self._alias_owner.get(key) == src:
                self._alias_owner[key] = dst
                if all(clean_name(a) != key for a in target.aliases):
                    target.aliases.append(alias)
            counts = self._alias_counts.get(key)
            if counts and src in counts:
                counts[dst] += counts.pop(src)
            by = self._alias_by.get(key)
            if by and src in by:
                moved = by.pop(src)
                kept = by.setdefault(dst, (set(), set(), set()))
                for a, b in zip(kept, moved):
                    a |= b
        for table in (self._ambiguous, self._seed_conflicts):
            for key in [k for k, v in table.items() if src in v]:
                table[key].discard(src)
                if table is self._seed_conflicts and self._alias_owner.get(key) != dst:
                    table[key].add(dst)
                elif table is self._ambiguous:
                    table[key].add(dst)
                if len(table[key]) < (2 if table is self._ambiguous else 1):
                    del table[key]
        target.atc |= entry.atc
        target.routes |= entry.routes
        target.otc |= entry.otc
        target.products += entry.products
        target.sources |= entry.sources
        if self._by_components.get(frozenset(entry.components)) == src:
            del self._by_components[frozenset(entry.components)]
        self._no_inn.discard(src)
        self._lexicon = None

    @classmethod
    def from_grls(cls, path: str | Path | Iterable[str | Path], **kwargs: Any) -> tuple["DrugCatalog", LoadReport]:
        catalog = cls()
        products, report = load_grls(path, **kwargs)
        catalog.add_products(products)
        return catalog, report

    # ------------------------------------------------------------------ словарь для нормализации
    def entries(self) -> list[LexEntry]:
        """Записи для Lexicon: стартовый набор первым, неоднозначные наименования исключены."""
        out = []
        for e in self._entries.values():
            aliases = tuple(a for a in e.aliases if self._usable(clean_name(a), e.canonical))
            out.append(LexEntry(e.canonical, aliases, e.components if len(e.components) > 1 else (),
                                e.classes, tuple(sorted(e.atc))))
        return out

    def mention_keys(self, exclude: Iterable[str] = ()) -> set[str]:
        """Ключи, которые можно искать в свободном тексте.

        Стартовый набор — все. ГРЛС — только наименования от 5 букв, не из стоп-списка
        и не совпадающие со словами других справочников (болезни, симптомы, аллергены).
        """
        blocked = {clean_name(x) for x in exclude} | MENTION_STOPWORDS
        keys = set()
        for key, owner in self._alias_owner.items():
            if not self._usable(key):
                continue
            if self._alias_source.get(key) == "seed":
                keys.add(key)
            elif (len(key.replace(" ", "")) >= 5 and key not in blocked
                  and key.split()[0] not in MENTION_PREPOSITIONS          # «от кашля», «для детей»
                  and not all(t in blocked for t in key.split())):
                keys.add(key)
        return keys

    def non_drug_words(self) -> set[str]:
        words: set[str] = set()
        for lex in (ref.CONDITIONS, ref.SYMPTOMS, ref.DIAGNOSES, ref.ALLERGENS):
            for e in lex:
                for a in (e.canonical, *e.aliases):
                    words.add(clean_name(a))
        return words

    def lexicon(self) -> Lexicon:
        if self._lexicon is None:
            self._lexicon = Lexicon(self.entries(), mention_keys=self.mention_keys(self.non_drug_words()))
        return self._lexicon

    # ------------------------------------------------------------------ сведения о препарате
    def __len__(self) -> int:
        return len(self._entries)

    def __contains__(self, canonical: str) -> bool:
        return canonical in self._entries

    def components(self, canonical: str) -> tuple[str, ...]:
        e = self._entries.get(canonical)
        return e.components if e else (canonical,)

    def atc_for(self, canonical: str, raw: str | None = None) -> tuple[str, ...]:
        """Коды АТХ: торгового наименования, если оно известно отдельно (местные формы), иначе МНН."""
        if raw:
            key = clean_name(raw)
            tokens = key.split()
            for n in range(len(tokens), 0, -1):
                codes = self._alias_atc.get(" ".join(tokens[:n]))
                if codes and self._alias_owner.get(" ".join(tokens[:n]), canonical) == canonical:
                    return tuple(sorted(codes))
        e = self._entries.get(canonical)
        return tuple(sorted(e.atc)) if e else ()

    def _alias_value(self, mapping: dict[str, set], canonical: str, raw: str | None) -> set | None:
        """Значение для торгового наименования (самый длинный префикс по словам), если оно известно."""
        if not raw:
            return None
        tokens = clean_name(raw).split()
        for n in range(len(tokens), 0, -1):
            key = " ".join(tokens[:n])
            if key in mapping and self._alias_owner.get(key, canonical) == canonical:
                return mapping[key]
        return None

    def routes_for(self, canonical: str, raw: str | None = None) -> tuple[str, ...]:
        found = self._alias_value(self._alias_routes, canonical, raw)
        if found is None:
            e = self._entries.get(canonical)
            found = e.routes if e else set()
        return tuple(r for r in ROUTES if r in found)

    def otc_for(self, canonical: str, raw: str | None = None) -> bool | None:
        """True — есть безрецептурная форма; False — только по рецепту; None — неизвестно."""
        found = self._alias_value(self._alias_otc, canonical, raw)
        if found is None:
            e = self._entries.get(canonical)
            found = e.otc if e else set()
        return True if True in found else (False if False in found else None)

    def components_for(self, canonical: str, raw: str | None = None) -> tuple[str, ...]:
        """Состав; для неоднозначного наименования — вещества из всех его составов в ГРЛС."""
        if raw:
            tokens = clean_name(raw).split()
            for n in range(len(tokens), 0, -1):
                key = " ".join(tokens[:n])
                if key in self._alias_union and self._resolved.get(key) == canonical:
                    return self._alias_union[key]
        return self.components(canonical)

    def _info(self, raw: str, canonical: str, method: str = "exact") -> DrugInfo:
        return DrugInfo(raw, canonical, self.components_for(canonical, raw), self.atc_for(canonical, raw), True,
                        method, self.routes_for(canonical, raw), self.otc_for(canonical, raw))

    def info(self, nr: NormResult) -> DrugInfo:
        if not nr.matched:
            return DrugInfo(nr.raw, nr.normalized, (nr.normalized,), (), False, nr.method)
        return self._info(nr.raw, nr.normalized, nr.method)

    def resolve(self, name: str) -> DrugInfo:
        return self.info(self.lexicon().lookup(name))

    def from_key(self, canonical: str, raw_names: Sequence[str] = ()) -> DrugInfo:
        """Факт памяти (ключ + как называл покупатель) → сведения о препарате."""
        raw = raw_names[0] if raw_names else canonical
        if canonical in self._entries:
            return self._info(raw, canonical)
        nr = self.lexicon().lookup(raw)  # ключ не из этого справочника: пробуем по названию
        return self.info(nr) if nr.matched else DrugInfo(raw, canonical, (canonical,), (), False, "unmatched")

    def mentions(self, text: str) -> list[DrugInfo]:
        """Препараты, названные в тексте; raw — найденное наименование без падежного окончания."""
        return [self.info(replace(nr, raw=nr.cleaned)) for nr in self.lexicon().find_mentions(text)]

    # ------------------------------------------------------------------ статистика и кэш
    def summary(self) -> dict[str, Any]:
        by_source = Counter()
        for e in self._entries.values():
            by_source["+".join(sorted(e.sources))] += 1
        with_atc = sum(1 for e in self._entries.values() if e.atc)
        return {
            "entries": len(self._entries),
            "entries_by_source": dict(by_source),
            "combinations": sum(1 for e in self._entries.values() if len(e.components) > 1),
            "with_atc": with_atc,
            "with_routes": sum(1 for e in self._entries.values() if e.routes),
            "otc": sum(1 for e in self._entries.values() if True in e.otc),
            "rx_only": sum(1 for e in self._entries.values() if e.otc == {False}),
            "aliases": len(self._alias_owner),
            "ambiguous_aliases": len(self._ambiguous),
            "ambiguous_resolved": len(self._resolved),
            "no_inn_merged": self.stats["no_inn_merged"],
            "products": self.stats["products"],
            "no_inn_products": self.stats["no_inn"],
            "seed_priority_conflicts": len(self._seed_conflicts),
            "merged_inn_variants": self.stats["merged_inn_variants"],
        }

    def ambiguous(self, limit: int = 20) -> list[tuple[str, list[str]]]:
        """Неразрешённые наименования (составы встречаются одинаково часто): (наименование, составы)."""
        items = sorted(((k, v) for k, v in self._ambiguous.items() if k not in self._resolved),
                       key=lambda kv: (-len(kv[1]), kv[0]))
        return [(k, sorted(v)) for k, v in items[:limit]]

    def resolved(self, limit: int = 20) -> list[tuple[str, str, list[tuple[str, int]]]]:
        """Разрешённые по частоте: (наименование, выбранный состав, [(состав, строк ГРЛС), …])."""
        out = []
        for k in sorted(self._resolved, key=lambda k: (-len(self._ambiguous.get(k, ())), k))[:limit]:
            counts = self._alias_counts.get(k, Counter())
            out.append((k, self._resolved[k],
                        sorted(((c, counts.get(c, 0)) for c in self._ambiguous.get(k, ())), key=lambda x: (-x[1], x[0]))))
        return out

    def seed_conflicts(self, limit: int = 50) -> list[tuple[str, str, list[str]]]:
        """Синонимы стартового набора, которым в ГРЛС соответствует другой состав: (синоним, наш ключ, по ГРЛС)."""
        items = sorted(self._seed_conflicts.items())[:limit]
        return [(k, self._alias_owner.get(k, ""), sorted(v)) for k, v in items]

    def save(self, path: str | Path) -> None:
        """Кэш справочника (gzip JSON): ГРЛС разбирается один раз, дальше загружается за секунды."""
        data = {
            "version": CATALOG_VERSION,
            "entries": [{"canonical": e.canonical, "components": list(e.components), "aliases": e.aliases,
                         "atc": sorted(e.atc), "classes": list(e.classes), "sources": sorted(e.sources),
                         "products": e.products, "routes": sorted(e.routes), "otc": sorted(e.otc)}
                        for e in self._entries.values()],
            "alias_source": self._alias_source,
            "alias_atc": {k: sorted(v) for k, v in self._alias_atc.items()},
            "alias_routes": {k: sorted(v) for k, v in self._alias_routes.items()},
            "alias_otc": {k: sorted(v) for k, v in self._alias_otc.items()},
            "ambiguous": {k: sorted(v) for k, v in self._ambiguous.items()},
            "no_inn": sorted(self._no_inn),
            "resolved": self._resolved,
            "alias_union": {k: list(v) for k, v in self._alias_union.items()},
            "alias_counts": {k: dict(v) for k, v in self._alias_counts.items() if k in self._ambiguous},
            "seed_conflicts": {k: sorted(v) for k, v in self._seed_conflicts.items()},
            "stats": dict(self.stats),
        }
        with gzip.open(path, "wt", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False)

    @classmethod
    def load(cls, path: str | Path) -> "DrugCatalog":
        with gzip.open(path, "rt", encoding="utf-8") as f:
            data = json.load(f)
        if data.get("version") != CATALOG_VERSION:
            raise ValueError(f"кэш справочника собран другой версией кода ({data.get('version')}, нужна "
                             f"{CATALOG_VERSION}): соберите справочник из выгрузки заново")
        cat = cls(seed=[], alias_atc={})
        for d in data["entries"]:
            e = _Entry(d["canonical"], tuple(d["components"]), list(d["aliases"]), set(d["atc"]),
                       tuple(d["classes"]), set(d["sources"]), d["products"], set(d.get("routes", ())),
                       set(d.get("otc", ())))
            cat._entries[e.canonical] = e
            cat._by_components.setdefault(frozenset(e.components), e.canonical)
            for a in e.aliases:
                cat._alias_owner.setdefault(clean_name(a), e.canonical)
        cat._alias_source = dict(data["alias_source"])
        cat._alias_atc = {k: set(v) for k, v in data["alias_atc"].items()}
        cat._alias_routes = {k: set(v) for k, v in data.get("alias_routes", {}).items()}
        cat._alias_otc = {k: set(v) for k, v in data.get("alias_otc", {}).items()}
        cat._ambiguous = {k: set(v) for k, v in data["ambiguous"].items()}
        cat._no_inn = set(data.get("no_inn", ()))
        cat._resolved = dict(data.get("resolved", {}))
        cat._alias_union = {k: tuple(v) for k, v in data.get("alias_union", {}).items()}
        cat._alias_counts = {k: Counter(v) for k, v in data.get("alias_counts", {}).items()}
        cat._seed_conflicts = {k: set(v) for k, v in data.get("seed_conflicts", {}).items()}
        cat.stats = Counter(data["stats"])
        return cat
